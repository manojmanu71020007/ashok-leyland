"""Excel export module for Chandapura Electric Bus Depot.

Implements complete 11-sheet workbook export matching Section 33 requirements:
1. Schedule Allocation
2. Fleet
3. Charging Stretches
4. SOC Calculations
5. Infeasible Schedules
6. Reserve Sensitivity
7. Gaps
8. Bus Reuse
9. Night Halt
10. Data Quality
11. Assumptions

Freezes headers and formats columns for clear operational review.
"""

from pathlib import Path
from typing import List, Dict, Optional, Any
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from models import (
    BusVehicle, ScheduleDuty, AllocationResult, InfeasibleReport,
    SensitivityReport, NightHaltSwapReport, DynamicSparePoolReport, HeadwayGapReport
)
from allocator import (
    AllocatorConfig, allocate_bus, generate_infeasible_schedules_report,
    calculate_reserve_sensitivity, calculate_target_distance, calculate_soc_requirement
)
from optimizer import (
    build_compatible_duty_chains, optimize_depot_allocation,
    evaluate_night_halt_and_swaps, calculate_dynamic_spare_pool,
    find_route_headway_gaps
)


def style_header_row(ws, row_idx=1, fill_color="1F497D", font_color="FFFFFF"):
    """Style header row with bold font, colored fill, and center alignment."""
    fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
    font = Font(name="Calibri", size=11, bold=True, color=font_color)
    align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    for cell in ws[row_idx]:
        cell.fill = fill
        cell.font = font
        cell.alignment = align
    ws.row_dimensions[row_idx].height = 28
    ws.freeze_panes = f"A{row_idx + 1}"


def auto_fit_columns(ws, max_width=45):
    """Adjust column widths dynamically based on content."""
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), max_width)


def generate_depot_excel_export(
    schedules: List[ScheduleDuty],
    buses: List[BusVehicle],
    trips_by_key: Optional[Dict[str, Any]] = None,
    output_path: Path = Path("Chandapura_Depot_Battery_Allocation.xlsx"),
    allow_unknown_inspection: bool = True,
    dead_km_factor: float = AllocatorConfig.DEAD_KM_FACTOR,
) -> Path:
    """Generate the complete 11-sheet Excel workbook export."""
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # Pre-run allocations and reports
    allocations: List[AllocationResult] = [
        allocate_bus(s.schedule_key, buses, schedules, current_soc=100.0, allow_unknown_inspection=allow_unknown_inspection)
        for s in schedules
    ]
    infeasible_rep = generate_infeasible_schedules_report(schedules, buses)
    sensitivity_rep = calculate_reserve_sensitivity(schedules, buses)
    nh_rep = evaluate_night_halt_and_swaps(schedules, buses, allow_unknown_inspection=allow_unknown_inspection)
    spare_rep = calculate_dynamic_spare_pool(schedules, buses)
    gap_rep = find_route_headway_gaps(schedules, buses, trips_by_key)
    chains = build_compatible_duty_chains(schedules)

    # =========================================================================
    # 1. Sheet: Schedule Allocation
    # =========================================================================
    ws1 = wb.create_sheet(title="Schedule Allocation")
    ws1.append([
        "Shift", "Schedule ID", "Canonical Key", "Route", "Category", "Origin", "Destination",
        "Departure", "Arrival", "Route Length (km)", "Dead km", "Actual km", "Charging Points",
        "First Stretch (km)", "Longest Calc Stretch (km)", "Form 4 Stretch (km)", "Target Distance (km)",
        "First Stretch Req SOC %", "Longest Stretch Req SOC %", "Current SOC %", "SOC Margin %",
        "Recommended BM", "REG NO", "Bus Category", "Bus Actual Range (km)", "Allocation Status", "Reason / Warnings"
    ])
    for s, a in zip(schedules, allocations):
        origin = s.stretches[0].legs[0].origin if s.stretches and s.stretches[0].legs else ""
        dest = s.stretches[-1].legs[-1].destination if s.stretches and s.stretches[-1].legs else ""
        bus_match = next((b for b in buses if b.bm_no == a.chosen_bm_no), None) if a.chosen_bm_no else None
        
        ws1.append([
            s.shift, s.schedule_id, s.schedule_key, s.route, s.category, origin, dest,
            s.departure_time_str or "--:--", s.arrival_time_str or "--:--",
            s.route_length_km, s.dead_km, s.actual_km, s.charging_event_count,
            s.first_stretch_km, s.longest_calculated_stretch_km, s.form4_single_charge_km,
            a.longest_stretch_target_km, a.first_stretch_required_soc_pct, a.longest_stretch_required_soc_pct,
            a.current_soc_pct, a.soc_margin_pct, a.chosen_bm_no or "None", a.chosen_reg_no or "None",
            bus_match.category if bus_match else "", bus_match.actual_range_km if bus_match else "",
            a.allocation_status, a.reason
        ])
    style_header_row(ws1, fill_color="1F497D")
    auto_fit_columns(ws1)

    # =========================================================================
    # 2. Sheet: Fleet
    # =========================================================================
    ws2 = wb.create_sheet(title="Fleet")
    ws2.append([
        "BM NO", "REG NO", "Actual Range Final (km)", "Category A/B/C", "Current SOC %",
        "Manual Confirmed SOC", "Operational Status", "Defect Flag", "Inspection Status",
        "Current Duty", "Next Available Time"
    ])
    for b in buses:
        ws2.append([
            b.bm_no, b.reg_no or "None", b.actual_range_km, b.category, b.current_soc_pct,
            "YES" if b.is_manually_confirmed else "NO", b.status.value,
            "DEFECT" if b.defect_flag else "OK", b.inspection_status.value,
            b.current_duty or "Available", format_minutes_to_time(b.next_available_min)
        ])
    style_header_row(ws2, fill_color="244062")
    auto_fit_columns(ws2)

    # =========================================================================
    # 3. Sheet: Charging Stretches
    # =========================================================================
    ws3 = wb.create_sheet(title="Charging Stretches")
    ws3.append([
        "Schedule Key", "Shift", "Schedule ID", "Route", "Stretch Index", "Route Length (km)",
        "Target Distance (km)", "Leg Count", "Is First Stretch", "Is Longest Stretch", "Leg Details"
    ])
    for s in schedules:
        for str_info in s.stretches:
            leg_desc = " -> ".join([f"{l.origin} to {l.destination}" for l in str_info.legs[:3]])
            if len(str_info.legs) > 3:
                leg_desc += f" (+{len(str_info.legs)-3} more legs)"
            ws3.append([
                s.schedule_key, s.shift, s.schedule_id, s.route, str_info.stretch_index,
                str_info.route_length_km, str_info.target_distance_km, len(str_info.legs),
                "YES" if str_info.is_first_stretch else "NO",
                "YES" if str_info.is_longest_stretch else "NO",
                leg_desc
            ])
    style_header_row(ws3, fill_color="366092")
    auto_fit_columns(ws3)

    # =========================================================================
    # 4. Sheet: SOC Calculations
    # =========================================================================
    ws4 = wb.create_sheet(title="SOC Calculations")
    ws4.append([
        "Schedule Key", "Route", "Reference BM", "Bus Range (km)", "Target Distance (km)",
        "Base Required SOC %", "Reserve %", "Required SOC With Reserve %", "Current SOC %",
        "Departure Eligible", "Projected SOC After Stretch %", "SOC Margin %", "Status"
    ])
    ref_bus = max(buses, key=lambda b: b.actual_range_km)
    for a in allocations:
        calc = calculate_soc_requirement(a.longest_stretch_target_km, ref_bus.actual_range_km, 0.0, 100.0)
        ws4.append([
            a.schedule_key, a.route, ref_bus.bm_no, ref_bus.actual_range_km,
            a.longest_stretch_target_km, calc.base_required_soc_pct, calc.reserve_pct,
            calc.required_soc_with_reserve_pct, 100.0, "YES" if calc.departure_soc_eligible else "NO",
            calc.projected_soc_after_stretch_pct, calc.soc_margin_pct, calc.feasibility_status.value
        ])
    style_header_row(ws4, fill_color="4F81BD")
    auto_fit_columns(ws4)

    # =========================================================================
    # 5. Sheet: Infeasible Schedules
    # =========================================================================
    ws5 = wb.create_sheet(title="Infeasible Schedules")
    ws5.append([
        "Schedule Key", "Schedule ID", "Shift", "Route", "Form 4 Stretch (km)", "Calculated Stretch (km)",
        "Selected Stretch (km)", "Target Distance (km)", "Highest Range BM", "Highest Bus Range (km)",
        "Base Required SOC %", "Reserve Adjusted SOC %", "Shortfall (km)", "Shortfall SOC %",
        "Feasibility Status", "Recommended Action"
    ])
    for row in infeasible_rep.rows:
        ws5.append([
            row.schedule_key, row.schedule_id, row.shift, row.route, row.form4_stretch_km,
            row.calculated_stretch_km, row.selected_stretch_km, row.target_distance_km,
            row.highest_range_bm, row.highest_bus_range_km, row.base_required_soc_pct,
            row.reserve_adjusted_soc_pct, row.shortfall_km, row.shortfall_soc_pct,
            row.status.value, row.recommended_action
        ])
    style_header_row(ws5, fill_color="C00000")
    auto_fit_columns(ws5)

    # =========================================================================
    # 6. Sheet: Reserve Sensitivity
    # =========================================================================
    ws6 = wb.create_sheet(title="Reserve Sensitivity")
    ws6.append([
        "Reserve %", "Physically Impossible Schedules", "Reserve-Infeasible Schedules",
        "Feasible Schedules", "Total Schedules", "Reference Bus BM", "Reference Bus Range (km)"
    ])
    for s_row in sensitivity_rep.table:
        ws6.append([
            f"{s_row.reserve_pct}%", s_row.physically_impossible_count,
            s_row.reserve_infeasible_count, s_row.feasible_count, s_row.total_schedules,
            sensitivity_rep.reference_bus_bm, sensitivity_rep.reference_bus_range_km
        ])
    style_header_row(ws6, fill_color="953735")
    auto_fit_columns(ws6)

    # =========================================================================
    # 7. Sheet: Gaps
    # =========================================================================
    ws7 = wb.create_sheet(title="Gaps")
    ws7.append([
        "Route", "Origin", "Previous Departure", "Next Departure", "Previous Schedule Key",
        "Next Schedule Key", "Headway Gap (Minutes)", "Threshold Flag (>60 min)",
        "Recommended Spare BM", "Recommended REG NO", "Available Candidates Count", "Notes"
    ])
    for g in gap_rep.gaps:
        ws7.append([
            g.route, g.origin, g.prev_departure_str, g.next_departure_str,
            g.prev_schedule_key, g.next_schedule_key, g.gap_minutes, "FLAGGED (>60 min)",
            g.recommended_bus or "None", g.recommended_reg or "None",
            len(g.candidate_spare_buses), g.notes
        ])
    style_header_row(ws7, fill_color="E26B00")
    auto_fit_columns(ws7)

    # =========================================================================
    # 8. Sheet: Bus Reuse
    # =========================================================================
    ws8 = wb.create_sheet(title="Bus Reuse")
    ws8.append([
        "Chain ID", "First Duty Key", "First Departure", "First Arrival",
        "Second Duty Key", "Second Departure", "Second Arrival", "Turnaround Duration (Minutes)",
        "Turnaround Charging Enabled", "Total Chained Target Distance (km)", "Reusability Assessment"
    ])
    for c in chains:
        if c.is_multi_duty and len(c.schedules) >= 2:
            s1, s2 = c.schedules[0], c.schedules[1]
            ws8.append([
                c.chain_id, s1.schedule_key, s1.departure_time_str or "--:--", s1.arrival_time_str or "--:--",
                s2.schedule_key, s2.departure_time_str or "--:--", s2.arrival_time_str or "--:--",
                c.turnaround_duration_min, "YES (Configurable)", c.total_target_distance_km,
                "Compatible duty chain with valid layover; requires turnaround charging for combined range."
            ])
    style_header_row(ws8, fill_color="60497A")
    auto_fit_columns(ws8)

    # =========================================================================
    # 9. Sheet: Night Halt
    # =========================================================================
    ws9 = wb.create_sheet(title="Night Halt")
    ws9.append([
        "Schedule ID", "Schedule Key", "Shift", "Route", "Fixed BM", "Fixed REG",
        "Actual BM", "Actual REG", "Recommended BM", "Recommended REG",
        "Out Time", "In Time", "Crosses Midnight", "Duration (Minutes)",
        "Allocation Mismatch", "Swap Required", "Current Assignment Status", "Reason"
    ])
    for r in nh_rep.nh_records:
        ws9.append([
            r.schedule_id, r.schedule_key or "", r.shift, r.route, r.fix_bm or "None", r.fix_reg or "None",
            r.actual_bm or "None", r.actual_reg or "None", r.recommended_bm or "None", r.recommended_reg or "None",
            r.out_time_str, r.in_time_str, "YES" if r.crosses_midnight else "NO", r.duration_min,
            "YES" if r.is_mismatch else "NO", "SWAP REQUIRED" if r.swap_required else "OK",
            r.current_assignment_status, r.reason
        ])
    for r in nh_rep.general_records:
        ws9.append([
            r.schedule_id, r.schedule_key or "", r.shift, r.route, r.fix_bm or "Unknown", r.fix_reg or "Unknown",
            r.actual_bm or "Unknown", r.actual_reg or "Unknown", r.recommended_bm or "None", r.recommended_reg or "None",
            r.out_time_str, r.in_time_str, "YES" if r.crosses_midnight else "NO", r.duration_min,
            "NO", "NO", r.current_assignment_status, r.reason
        ])
    style_header_row(ws9, fill_color="31859B")
    auto_fit_columns(ws9)

    # =========================================================================
    # 10. Sheet: Data Quality
    # =========================================================================
    ws10 = wb.create_sheet(title="Data Quality")
    ws10.append([
        "Item ID", "Source File / Component", "Audit Finding", "Operational Impact", "System Treatment / Resolution"
    ])
    audit_items = [
        ("DQ-01", "OHM_004-_Check_sheet_for_inspection.xlsx", "Workbook contains blank operational procedure templates/KPI monitoring sheets without bus-level defect logs.", "Cannot infer that all 122 buses are defect-free.", "Defaulted all buses to UNKNOWN_NO_RECORD; added ALLOW_UNKNOWN_INSPECTION_FOR_ALLOCATION toggle."),
        ("DQ-02", "SCH_DATA_NH_new__1_.xlsx (SCH DATA NH)", "Master sheet contains 82 duties (71 Night Halt + 10 Shift A + 1 Shift B). Exactly 39 General-shift duties are missing.", "Cannot establish actual incumbent bus for 39 General-shift duties.", "Explicitly marked Current Assignment = Unknown; prohibited guessing from route or timing."),
        ("DQ-03", "Form 4 Shift A & Shift B Workbooks", "10 schedule IDs (356Z/6, 356Z/7, 600F/53, etc.) exist identically in both Shift A and Shift B.", "Global lookups by schedule_id alone collide and produce ambiguous results.", "Implemented canonical Schedule Key (Shift|Schedule ID); allocate_bus without shift returns AMBIGUOUS_SCHEDULE_ID."),
        ("DQ-04", "Form 4 Duty-Card Leg Origin Fields", "Charging events logged with various cases ('CHARGING', 'Charging', 'CHARGING ').", "Exact string matching misses mid-route battery charging opportunities.", "Implemented case-insensitive, whitespace-tolerant charging detection across all Form 4 files."),
        ("DQ-05", "Night Halt Duty Cards", "Duties depart between 12:30-18:00 and arrive between 06:00-12:00 next day.", "Clock-time subtraction produces negative durations.", "Midnight day normalization: Arrival clock < Departure clock triggers Day + 1 calendar offset."),
        ("DQ-06", "Consolidation Single Charge vs Calculated", "86.2% of Form 4 Single-Charge distances match Rule C calculated stretches within 2.0 km.", "Form 4 values occasionally omit overnight boundary splits.", "Supported dual stretch sources (FORM4 and CALCULATED) with consistency status classification.")
    ]
    for item in audit_items:
        ws10.append(list(item))
    style_header_row(ws10, fill_color="595959")
    auto_fit_columns(ws10)

    # =========================================================================
    # 11. Sheet: Assumptions
    # =========================================================================
    ws11 = wb.create_sheet(title="Assumptions")
    ws11.append([
        "Assumption ID", "Parameter / Policy", "Configured Default Value", "Operational Rationale & Dataset Evidence"
    ])
    assumptions_items = [
        ("ASM-01", "Dead-KM Allowance Factor", "1.10 (10% dead-km addition)", "Derived from Chandapura Consolidation Form 4 Dead KM column average ratio (Route Length * 1.10 = Actual KM)."),
        ("ASM-02", "Minimum Turnaround Duration", "20 Minutes", "Minimum required buffer between arrival of Duty A and departure of Duty B for crew/bus changeover."),
        ("ASM-03", "Turnaround Charging Duration", "45 Minutes", "Standard fast-charging layover window required to replenish battery between consecutive chained shifts."),
        ("ASM-04", "Target SOC After Charge", "100.0%", "Default theoretical battery level after dedicated charging stop (operator manual update overrides this)."),
        ("ASM-05", "Reserve SOC Buffer", "0.0% (Sensitivity tested at 0%, 5%, 10%, 15%)", "Baseline calculation uses 0% reserve; sensitivity report answers how many routes remain feasible at higher buffers."),
        ("ASM-06", "Headway Gap Flag Threshold", "60 Minutes", "Service intervals between consecutive trips on the same route/origin exceeding 60 min indicate headway deficits."),
        ("ASM-07", "Category Matching Hierarchy", "Complex -> Cat A; Moderate -> Cat B; Standard -> Cat C/B", "Conserves high-range Category A buses for high-demand routes, avoiding early fleet range exhaustion.")
    ]
    for item in assumptions_items:
        ws11.append(list(item))
    style_header_row(ws11, fill_color="376092")
    auto_fit_columns(ws11)

    # Save workbook
    wb.save(output_path)
    return output_path


def format_minutes_to_time(minutes: Optional[int]) -> str:
    """Format minutes from midnight to HH:MM format."""
    if minutes is None:
        return "--:--"
    m = minutes % 1440
    return f"{m // 60:02d}:{m % 60:02d}"
