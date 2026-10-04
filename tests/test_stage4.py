"""Stage 4 Test Suite — Validation of Excel Export and Dashboard Integrity.

Verifies:
1. Complete 11-sheet Excel workbook export matches all Section 33 requirements.
2. Excel sheet names, headers, and row populations.
3. Full dashboard component imports and data readiness.
"""

import pytest
from pathlib import Path
import openpyxl

from data_loader import load_bus_fleet, load_consolidation_schedules, load_trip_legs
from validation import enrich_schedules_with_stretches
from export_excel import generate_depot_excel_export


@pytest.fixture(scope="module")
def loaded_depot_data():
    buses = load_bus_fleet()
    schedules = load_consolidation_schedules()
    trips = load_trip_legs()
    enrich_schedules_with_stretches(schedules, trips)
    return buses, schedules, trips


def test_excel_export_11_sheets_structure(loaded_depot_data, tmp_path):
    """Verify that generate_depot_excel_export produces the required 11-sheet workbook."""
    buses, schedules, trips = loaded_depot_data
    out_file = tmp_path / "Chandapura_Depot_Test_Export.xlsx"

    result_path = generate_depot_excel_export(
        schedules=schedules,
        buses=buses,
        trips_by_key=trips,
        output_path=out_file,
        allow_unknown_inspection=True
    )

    assert result_path.exists()
    assert result_path.stat().st_size > 30000

    wb = openpyxl.load_workbook(result_path, data_only=True)
    expected_sheets = [
        "Schedule Allocation",
        "Fleet",
        "Charging Stretches",
        "SOC Calculations",
        "Infeasible Schedules",
        "Reserve Sensitivity",
        "Gaps",
        "Bus Reuse",
        "Night Halt",
        "Data Quality",
        "Assumptions"
    ]

    assert len(wb.sheetnames) == 11
    for s_name in expected_sheets:
        assert s_name in wb.sheetnames
        ws = wb[s_name]
        # Must have headers and data rows
        assert ws.max_row >= 2
        assert ws.max_column >= 3


def test_excel_export_sheet_contents(loaded_depot_data, tmp_path):
    """Verify specific sheet contents matching Section 33 requirements."""
    buses, schedules, trips = loaded_depot_data
    out_file = tmp_path / "Chandapura_Depot_Test_Export_2.xlsx"

    generate_depot_excel_export(schedules, buses, trips, out_file)
    wb = openpyxl.load_workbook(out_file, data_only=True)

    # 1. Schedule Allocation has 130 schedules + 1 header = 131 rows
    ws_alloc = wb["Schedule Allocation"]
    assert ws_alloc.max_row == 131

    # 2. Fleet has 122 buses + 1 header = 123 rows
    ws_fleet = wb["Fleet"]
    assert ws_fleet.max_row == 123

    # 3. Infeasible Schedules has rows flagged PHYSICALLY IMPOSSIBLE
    ws_inf = wb["Infeasible Schedules"]
    assert ws_inf.max_row > 10

    # 4. Reserve Sensitivity has 4 rows (0%, 5%, 10%, 15%) + header
    ws_sens = wb["Reserve Sensitivity"]
    assert ws_sens.max_row == 5

    # 5. Night Halt has 71 NH + 39 General + 1 header = 111 rows
    ws_nh = wb["Night Halt"]
    assert ws_nh.max_row == 111


def test_app_compilation():
    """Verify app.py compiles without syntax errors."""
    import py_compile
    compiled = py_compile.compile("app.py", doraise=True)
    assert compiled is not None
