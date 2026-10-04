"""Data loading and Excel parsing utilities for Chandapura EV depot.

Handles:
- Case-insensitive, whitespace-tolerant charging detection
- Robust Excel time parsing (numeric fraction of day, datetime.time, string HH:MM)
- Midnight crossing detection
- Shift-qualified canonical schedule keys (f"{shift}|{schedule_id}")
- BM NO to REG NO mapping from master sheets
- 3-state inspection data audit (CONFIRMED_OK, DEFECT, UNKNOWN_NO_RECORD)
"""

from pathlib import Path
from typing import List, Dict, Tuple, Optional, Any
import datetime
import openpyxl

from models import BusVehicle, BusStatus, InspectionStatus, ScheduleDuty, TripLeg


def normalize_shift(shift_str: Optional[str]) -> str:
    """Normalize shift string for deterministic cross-file matching."""
    if not shift_str:
        return "UNKNOWN"
    s = str(shift_str).strip().lower()
    if "night" in s or "nh" in s or "hault" in s:
        return "Night Halt"
    if "gen" in s:
        return "General shift"
    if s == "a" or "shift a" in s:
        return "Shift A"
    if s == "b" or "shift b" in s:
        return "Shift B"
    return shift_str.strip()


def parse_excel_time(val: Any) -> Optional[int]:
    """Parse Excel time values into minutes from midnight (0..1439).
    
    Accepts:
    - datetime.time / datetime.datetime
    - float / int (Excel fraction of 24h day)
    - str ('HH:MM' or 'HH:MM:SS')
    """
    if val is None:
        return None
    if isinstance(val, (datetime.time, datetime.datetime)):
        return val.hour * 60 + val.minute
    if isinstance(val, (int, float)):
        # Excel fraction of day: 1.0 = 24h = 1440 min
        mins = round(float(val) * 24.0 * 60.0)
        return mins % 1440
    if isinstance(val, str):
        s = val.strip()
        parts = s.split(":")
        if len(parts) >= 2:
            try:
                hour = int(parts[0])
                minute = int(parts[1])
                return (hour * 60 + minute) % 1440
            except ValueError:
                pass
    return None


def format_minutes_to_time(minutes: Optional[int]) -> str:
    """Format minutes from midnight to HH:MM format."""
    if minutes is None:
        return "--:--"
    m = minutes % 1440
    return f"{m // 60:02d}:{m % 60:02d}"


def is_charging_event(origin_val: Any, dest_val: Any = None, chg_val: Any = None) -> bool:
    """Case-insensitive and whitespace-tolerant matching for charging events.
    
    Recognizes:
    - 'CHARGING', 'Charging', 'charging', 'CHARGING ' in Origin
    - Also tolerates 'charging' in Destination or Dedicated Charging column.
    """
    orig_str = str(origin_val).strip().lower() if origin_val is not None else ""
    if orig_str == "charging" or "charg" in orig_str:
        return True
    if dest_val is not None:
        dest_str = str(dest_val).strip().lower()
        if dest_str == "charging":
            return True
    if chg_val is not None:
        c_str = str(chg_val).strip().lower()
        if c_str == "charging":
            return True
    return False


def load_bus_fleet(
    consolidation_path: Path = Path("data/Form_4_Consolidation_-_Chandapura.xlsx"),
    veh_master_path: Path = Path("data/SCH_DATA_NH_new__1_.xlsx")
) -> List[BusVehicle]:
    """Load the full 122-bus fleet with degradation actual range and REG NO mapping.
    
    Since the inspection workbook contains only blank templates, all buses
    initially receive inspection_status = UNKNOWN_NO_RECORD rather than CONFIRMED_OK.
    """
    # 1. Load REG NO mapping from VEH.NO sheet
    reg_map: Dict[str, str] = {}
    if veh_master_path.exists():
        wb_veh = openpyxl.load_workbook(veh_master_path, data_only=True)
        if "VEH.NO" in wb_veh.sheetnames:
            ws_v = wb_veh["VEH.NO"]
            for row in list(ws_v.iter_rows(values_only=True))[1:]:
                if row and len(row) >= 3 and row[1]:
                    bm = str(row[1]).strip().upper()
                    reg = str(row[2]).strip() if row[2] else None
                    if reg:
                        reg_map[bm] = reg

    # 2. Load bus categories and ranges
    wb_cons = openpyxl.load_workbook(consolidation_path, data_only=True)
    ws_bus = wb_cons["Bus category"]
    buses: List[BusVehicle] = []
    
    for row in list(ws_bus.iter_rows(values_only=True))[1:]:
        if not row or not row[1]:
            continue
        bm_no = str(row[1]).strip().upper()
        range_val = float(row[2]) if row[2] is not None else 0.0
        cat_val = str(row[3]).strip().upper() if row[3] is not None else "C"
        reg_no = reg_map.get(bm_no)
        
        buses.append(BusVehicle(
            bm_no=bm_no,
            reg_no=reg_no,
            actual_range_km=round(range_val, 4),
            category=cat_val,
            current_soc_pct=100.0,
            status=BusStatus.AVAILABLE,
            inspection_status=InspectionStatus.UNKNOWN_NO_RECORD,
            defect_flag=False
        ))
    
    # Sort deterministically by BM NO
    buses.sort(key=lambda b: b.bm_no)
    return buses


def load_consolidation_schedules(
    consolidation_path: Path = Path("data/Form_4_Consolidation_-_Chandapura.xlsx")
) -> List[ScheduleDuty]:
    """Load 130 consolidation schedules with canonical shift-aware Schedule Keys."""
    wb = openpyxl.load_workbook(consolidation_path, data_only=True)
    ws = wb["Schedule category"]
    rows = list(ws.iter_rows(values_only=True))
    schedules: List[ScheduleDuty] = []
    
    for r in rows[1:]:
        if not r or r[1] is None:
            continue
        shift = normalize_shift(r[0])
        sched_id = str(r[1]).strip()
        route_len = float(r[2]) if r[2] is not None else 0.0
        dead_km = float(r[8]) if r[8] is not None else 0.0
        actual_km = float(r[9]) if r[9] is not None else 0.0
        single_charge = float(r[10]) if r[10] is not None else None
        cat = str(r[11]).strip() if r[11] is not None else "Standard"
        remarks = str(r[12]).strip() if r[12] is not None else None
        
        # Route is prefix before slash
        route = sched_id.split("/")[0] if "/" in sched_id else sched_id
        canonical_key = f"{shift}|{sched_id}"
        
        warning_flags = []
        if single_charge is None:
            warning_flags.append("MISSING FORM4 STRETCH")
            
        schedules.append(ScheduleDuty(
            schedule_key=canonical_key,
            shift=shift,
            schedule_id=sched_id,
            route=route,
            route_length_km=round(route_len, 2),
            dead_km=round(dead_km, 2),
            actual_km=round(actual_km, 2),
            form4_single_charge_km=single_charge,
            category=cat,
            remarks=remarks,
            warning_flags=warning_flags
        ))
        
    return schedules


def load_trip_legs(
    data_dir: Path = Path("data")
) -> Dict[str, List[TripLeg]]:
    """Load trip-level duty card legs from all Form 4 files.
    
    Returns:
    Dict mapping canonical schedule_key ("{shift}|{schedule_id}") -> List[TripLeg]
    """
    file_configs = [
        ("Shift A", data_dir / "Form_4_Shift_A.xlsx"),
        ("Shift B", data_dir / "Form_4_Shift_B.xlsx"),
        ("General shift", data_dir / "Form_4_General.xlsx"),
        ("Night Halt", data_dir / "Form_4_Night_Hault.xlsx"),
    ]
    
    trips_by_key: Dict[str, List[TripLeg]] = {}
    
    for default_shift, file_path in file_configs:
        if not file_path.exists():
            continue
        wb = openpyxl.load_workbook(file_path, data_only=True)
        ws = wb["Sheet1"]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
            
        headers = [str(c).strip() if c is not None else "" for c in rows[0]]
        sched_idx = headers.index("Schedule") if "Schedule" in headers else -1
        orig_idx = headers.index("Origin") if "Origin" in headers else -1
        dest_idx = headers.index("Destination") if "Destination" in headers else -1
        dep_idx = headers.index("Departure") if "Departure" in headers else -1
        arr_idx = headers.index("Arrival") if "Arrival" in headers else -1
        len_idx = headers.index("Route Length") if "Route Length" in headers else -1
        route_idx = headers.index("Route") if "Route" in headers else -1
        shift_idx = headers.index("Shift") if "Shift" in headers else -1
        rest_idx = headers.index("Rest") if "Rest" in headers else -1
        chg_idx = headers.index("Charging") if "Charging" in headers else -1
        
        current_sched = None
        current_route = ""
        current_shift = default_shift
        
        for r_idx, r in enumerate(rows[1:], 2):
            if not any(r):
                continue
                
            if sched_idx >= 0 and r[sched_idx] is not None and str(r[sched_idx]).strip() != "":
                current_sched = str(r[sched_idx]).strip()
            if not current_sched:
                continue
                
            if route_idx >= 0 and r[route_idx] is not None and str(r[route_idx]).strip() != "":
                current_route = str(r[route_idx]).strip()
            elif "/" in current_sched:
                current_route = current_sched.split("/")[0]
                
            if shift_idx >= 0 and r[shift_idx] is not None and str(r[shift_idx]).strip() != "":
                current_shift = normalize_shift(str(r[shift_idx]))
            else:
                current_shift = default_shift
                
            orig_raw = r[orig_idx] if orig_idx >= 0 else ""
            dest_raw = r[dest_idx] if dest_idx >= 0 else ""
            chg_raw = r[chg_idx] if chg_idx >= 0 else None
            
            orig = str(orig_raw).strip() if orig_raw is not None else ""
            dest = str(dest_raw).strip() if dest_raw is not None else ""
            
            is_charging = is_charging_event(orig, dest, chg_raw)
            
            raw_dep = r[dep_idx] if dep_idx >= 0 else None
            raw_arr = r[arr_idx] if arr_idx >= 0 else None
            dep_min = parse_excel_time(raw_dep)
            arr_min = parse_excel_time(raw_arr)
            
            crosses_midnight = False
            if dep_min is not None and arr_min is not None and arr_min < dep_min:
                crosses_midnight = True
                
            rlen = float(r[len_idx]) if len_idx >= 0 and r[len_idx] is not None else 0.0
            rest_raw = r[rest_idx] if rest_idx >= 0 else None
            rest_str = str(rest_raw).strip() if rest_raw is not None else None
            
            leg = TripLeg(
                sl_no=r_idx,
                route=current_route,
                schedule=current_sched,
                shift=current_shift,
                origin=orig,
                destination=dest,
                departure_min=dep_min,
                arrival_min=arr_min,
                crosses_midnight=crosses_midnight,
                route_length_km=round(rlen, 2),
                rest_str=rest_str,
                is_charging_event=is_charging,
                raw_departure=str(raw_dep) if raw_dep is not None else None,
                raw_arrival=str(raw_arr) if raw_arr is not None else None,
            )
            
            canonical_key = f"{current_shift}|{current_sched}"
            if canonical_key not in trips_by_key:
                trips_by_key[canonical_key] = []
            trips_by_key[canonical_key].append(leg)
            
    return trips_by_key


def load_inspection_audit(
    inspection_path: Path = Path("data/OHM_004-_Check_sheet_for_inspection.xlsx")
) -> Dict[str, Any]:
    """Audit the OHM_004 inspection workbook.
    
    Establishes UNKNOWN_NO_RECORD state due to blank operational templates.
    """
    if not inspection_path.exists():
        return {
            "status": InspectionStatus.UNKNOWN_NO_RECORD.value,
            "message": "Inspection file not found",
            "defects_by_bm": {}
        }
        
    wb = openpyxl.load_workbook(inspection_path, data_only=True)
    return {
        "status": InspectionStatus.UNKNOWN_NO_RECORD.value,
        "sheets": wb.sheetnames,
        "message": (
            "Workbook contains blank operational procedure templates/KPI monitoring sheets. "
            "No bus-level inspection records or defects are logged in the raw data. "
            "All buses are classified as UNKNOWN_NO_RECORD. Do not infer 'CONFIRMED_OK'."
        ),
        "defects_by_bm": {}
    }


def load_night_halt_roster(
    nh_path: Path = Path("data/SCH_DATA_NH_new__1_.xlsx")
) -> Dict[str, Dict[str, Any]]:
    """Load Night Halt and Shift A schedule vehicle commitments from SCH DATA NH."""
    if not nh_path.exists():
        return {}
        
    wb = openpyxl.load_workbook(nh_path, data_only=True)
    ws = wb["SCH DATA NH"]
    rows = list(ws.iter_rows(values_only=True))
    roster: Dict[str, Dict[str, Any]] = {}
    
    # Header is at row index 2 (row 3 in Excel)
    for r in rows[3:]:
        if not r or len(r) < 9 or not r[1]:
            continue
        sched_id = str(r[1]).strip()
        fix_bm = str(r[2]).strip().upper() if r[2] else None
        fix_reg = str(r[3]).strip() if r[3] else None
        swap_info = str(r[4]).strip() if r[4] else None
        bm_no = str(r[5]).strip().upper() if r[5] else None
        reg_no = str(r[6]).strip() if r[6] else None
        out_time = parse_excel_time(r[7])
        in_time = parse_excel_time(r[8])
        
        roster[sched_id] = {
            "schedule_id": sched_id,
            "fix_bm": fix_bm,
            "fix_reg": fix_reg,
            "swap_info": swap_info,
            "bm_no": bm_no,
            "reg_no": reg_no,
            "out_time_min": out_time,
            "in_time_min": in_time,
            "is_swap": bool(swap_info or (fix_bm and bm_no and fix_bm != bm_no))
        }
    return roster
