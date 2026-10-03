import openpyxl
import os
import json
import csv

excel_path = r"C:\Users\Manoj K\.gemini\antigravity\brain\24a37357-21e1-481b-8cb2-ae0889c5386e\.user_uploaded\media_1791004754246.xlsx"
wb = openpyxl.load_workbook(excel_path, data_only=True)

# 1. Parse Bus Categories
ws_bus = wb['Bus category']
bus_data = []
for r in list(ws_bus.iter_rows(values_only=True))[1:]:
    if any(r) and r[1]:
        sl, bm, rng, cat = r[:4]
        bus_data.append({
            'sl_no': int(sl) if sl is not None else len(bus_data) + 1,
            'bm_number': str(bm).strip(),
            'actual_range_km': round(float(rng), 4),
            'category': str(cat).strip()
        })

# 2. Parse Schedule Categories
ws_sched = wb['Schedule category']
sched_data = []
for r in list(ws_sched.iter_rows(values_only=True))[1:]:
    if any(r) and r[1]:
        shift, sched_id, r_len, n_rest, cum_t, chg_occ, no_chg, rl_sched, dead_km, act_km, single_chg, cat, remarks = r[:13]
        
        # format cum_time
        if hasattr(cum_t, 'strftime'):
            cum_str = cum_t.strftime('%H:%M:%S')
        else:
            cum_str = str(cum_t) if cum_t else ''
            
        sched_data.append({
            'shift': str(shift).strip() if shift else '',
            'schedule_id': str(sched_id).strip(),
            'route': str(sched_id).strip().split('/')[0],
            'route_length_km': round(float(r_len), 2) if r_len is not None else 0.0,
            'no_of_rest': int(n_rest) if n_rest is not None else 0,
            'cum_time': cum_str,
            'charging_occ': int(chg_occ) if chg_occ is not None and str(chg_occ).isdigit() else (1 if str(chg_occ).strip() == '1' else 0),
            'no_charging_occ_per_sched': int(no_chg) if no_chg is not None and str(no_chg).isdigit() else 0,
            'route_length_per_sched_km': round(float(rl_sched), 2) if rl_sched is not None else 0.0,
            'dead_km': round(float(dead_km), 2) if dead_km is not None else 0.0,
            'actual_km': round(float(act_km), 2) if act_km is not None else 0.0,
            'single_charge_req_km': round(float(single_chg), 2) if single_chg is not None else 0.0,
            'category': str(cat).strip() if cat else '',
            'remarks': str(remarks).strip() if remarks else ''
        })

# Save to vehicles directory
os.makedirs("vehicles", exist_ok=True)

with open("vehicles/bus_categories.json", "w", encoding="utf-8") as f:
    json.dump(bus_data, f, indent=2)

with open("vehicles/schedule_categories.json", "w", encoding="utf-8") as f:
    json.dump(sched_data, f, indent=2)

# Also CSV for easy inspection
with open("vehicles/bus_categories.csv", "w", newline="", encoding="utf-8") as f:
    writer = csv.DictWriter(f, fieldnames=['sl_no', 'bm_number', 'actual_range_km', 'category'])
    writer.writeheader()
    writer.writerows(bus_data)

with open("vehicles/schedule_categories.csv", "w", newline="", encoding="utf-8") as f:
    writer = csv.DictWriter(f, fieldnames=[
        'shift', 'schedule_id', 'route', 'route_length_km', 'no_of_rest', 'cum_time', 
        'charging_occ', 'no_charging_occ_per_sched', 'route_length_per_sched_km', 
        'dead_km', 'actual_km', 'single_charge_req_km', 'category', 'remarks'
    ])
    writer.writeheader()
    writer.writerows(sched_data)

print(f"Successfully saved {len(bus_data)} bus categories and {len(sched_data)} schedule categories.")
